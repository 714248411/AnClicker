"""Offline legacy migration: independent branches, complete report, no execution.

Usage: python migration_tool.py old.xlsx --output output-parent
Only successful, fully validated branches become executable workbooks. A copy
of the original always accompanies the report, including unsupported content.
"""
from __future__ import annotations

import argparse
import ast
import hashlib
import json
import ntpath
import re
import shutil
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook

from graph_repository import GraphRepository, SETTINGS_SHEET_HEADERS, WorkbookValidationError, GraphValidationError
from legacy_workbook import (
    LEGACY_HEADERS, DATABASE_HEADERS, convert_legacy_workbook, is_current_workbook,
)


@dataclass
class MigrationPlan:
    workbooks: dict = field(default_factory=dict)
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    branches: dict = field(default_factory=dict)
    default_branch: str = ""
    source_commands: int = 0
    missing_resources: list = field(default_factory=list)

    @property
    def converted_commands(self):
        return sum(row[1] != '报错跳转' for book in self.workbooks.values()
                   for row in book['命令'].iter_rows(min_row=2, values_only=True))

    def close(self):
        for book in self.workbooks.values():
            book.close()


def _read_settings(source):
    """Read both five-column settings and pre-SQLite two-column INI exports."""
    rows, sections, branch_data = [], {}, {}
    if '设置' not in source.sheetnames:
        return rows, sections, branch_data
    sheet = source['设置']
    header = list(next(sheet.values))
    if header == SETTINGS_SHEET_HEADERS:
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if all(value is None for value in row):
                continue
            if row[0] == '分支':
                branch_data[str(row[1])] = {'shortcut': row[2] or '', 'repeat': row[3] if row[3] is not None else 1}
            else:
                rows.append(list(row))
        return rows, sections, branch_data
    if sheet.max_column != 2:
        raise WorkbookValidationError('设置表既不是五列设置，也不是两列 INI 导出格式')
    section = None
    for index, (key, value) in enumerate(sheet.values, 1):
        if key is None and value is None:
            continue
        if isinstance(key, str) and key.startswith('[') and key.endswith(']') and value is None:
            section = key[1:-1]
            if section in sections:
                raise WorkbookValidationError(f'设置表第 {index} 行：重复的设置分组')
            sections[section] = {}
        elif section is None or not isinstance(key, str) or key in sections[section]:
            raise WorkbookValidationError(f'设置表第 {index} 行：缺少分组或设置名称重复')
        else:
            sections[section][key] = value
    config = sections.get('Config', {})
    # Do not silently change application-wide hotkeys, timing or UI preferences.
    # Retain the entire old configuration as inert data and describe it in report.
    if '图像匹配精度' in config:
        rows.append(['设置', '图像匹配精度', config['图像匹配精度'], None, None])
    resources = []
    for value in sections.get('资源文件夹路径', {}).values():
        if not isinstance(value, str) or not value.strip():
            raise WorkbookValidationError('旧资源文件夹路径不是有效文本')
        value = ntpath.normpath(value) if ntpath.splitdrive(value)[0] else value
        if value not in resources:
            resources.append(value)
            rows.append(['资源文件夹', value, None, None, len(resources) - 1])
    for name, raw in sections.get('分支', {}).items():
        try:
            value = ast.literal_eval(raw)
            if not isinstance(value, (tuple, list)) or len(value) != 2:
                raise ValueError('需要快捷键和次数两个值')
            shortcut, repeats = value
            if (not isinstance(shortcut, str) or isinstance(repeats, bool)
                    or not isinstance(repeats, int) or (repeats != -1 and repeats < 1)):
                raise ValueError('快捷键或次数无效')
            branch_data[name] = {'shortcut': shortcut, 'repeat': repeats}
        except (ValueError, TypeError, SyntaxError) as error:
            raise WorkbookValidationError(f'旧分支“{name}”设置无法安全解析：{error}') from error
    return rows, sections, branch_data


def inspect_migration(source, base_directory=None):
    """Scan every branch. Do not mutate source, caller DB, or run instructions."""
    plan = MigrationPlan()
    try:
        settings_rows, sections, plan.branches = _read_settings(source)
    except WorkbookValidationError as error:
        plan.errors.append({'sheet': '设置', 'error': str(error)})
        return plan
    config = sections.get('Config', {}) or {
        row[1]: row[2] for row in settings_rows if row[0] == '设置'
    }
    plan.default_branch = str(config.get('当前分支', ''))
    if sections:
        plan.warnings.append('旧窗口大小、全局/分支快捷键、三方接口及其他旧配置已完整归档，不自动覆盖新版应用偏好；旧全局时间间隔/持续时间/暂停时间未自动套用。')
    elif any(branch['shortcut'] for branch in plan.branches.values()):
        plan.warnings.append('旧分支快捷键已归档，不注册为新版全局快捷键。')
    resources = [row[1] for row in settings_rows if row[0] == '资源文件夹']
    branch_targets = {sheet.title: (_branch_filename(index, sheet.title),
                      sum(any(v is not None for v in row) for row in sheet.iter_rows(min_row=2, values_only=True)))
                      for index, sheet in enumerate((s for s in source if s.title != '设置'), 1)}
    # A private scratch repository validates the real protocol and graph invariants.
    with tempfile.TemporaryDirectory(prefix='anclicker-migrate-check-') as temporary:
        from 数据库操作 import DatabaseOperation
        path = str(Path(temporary) / 'check.db')
        DatabaseOperation(path)
        repository = GraphRepository(path)
        for sheet in source:
            if sheet.title == '设置':
                continue
            header = list(next(sheet.values))
            values = [(index, list(row)) for index, row in enumerate(
                sheet.iter_rows(min_row=2, values_only=True), 2
            ) if any(value is not None for value in row)]
            plan.source_commands += len(values)
            if header not in (LEGACY_HEADERS, DATABASE_HEADERS,
                              LEGACY_HEADERS + ['隶属分支'], DATABASE_HEADERS + ['隶属分支']):
                plan.errors.append({'sheet': sheet.title, 'error': '未识别的指令表头，原数据已保留'})
                continue
            if not values:
                plan.warnings.append(f'空分支“{sheet.title}”保留在原文件中，不生成可执行任务。')
                continue
            split = Workbook()
            split.active.title = sheet.title
            split.active.append(header)
            for _, row in values:
                split.active.append(row)
            settings = split.create_sheet('设置')
            settings.append(SETTINGS_SHEET_HEADERS)
            for row in settings_rows:
                # Paths/default branch are assigned to each new output instead.
                if row[0] == '设置' and row[1] in {'当前文件路径', '当前分支', '运行重复次数'}:
                    continue
                settings.append(row)
            branch = plan.branches.get(sheet.title, {'shortcut': '', 'repeat': 1})
            repeats = branch['repeat']
            if isinstance(repeats, bool) or not isinstance(repeats, int) or (repeats != -1 and not 1 <= repeats <= 999999):
                plan.errors.append({'sheet': sheet.title, 'error': '整组重复次数无效或超出范围'})
                split.close()
                continue
            settings.append(['设置', '运行重复次数', str(repeats), None, None])
            settings.append(['设置', '旧版迁移来源分支', sheet.title, None, None])
            settings.append(['设置', '旧版迁移配置归档', json.dumps(
                {'sections': sections, 'branch': branch}, ensure_ascii=False), None, None])
            try:
                converted = convert_legacy_workbook(split, base_directory, branch_targets=branch_targets)
                confidence = float(config.get('图像匹配精度', '0.8'))
                if not 0.01 <= confidence <= 1:
                    raise WorkbookValidationError('旧全局图像匹配精度无效')
                for row in converted['命令'].iter_rows(min_row=2):
                    parameters = json.loads(row[2].value)
                    if row[1].value in {'图像点击', '多图点击', '图像等待'}:
                        parameters.setdefault('精度', confidence)
                    if '图像路径' in parameters:
                        paths = str(parameters['图像路径']).splitlines()
                        resolved = []
                        for old_path in paths:
                            candidate = Path(old_path)
                            alternatives = [candidate]
                            if base_directory:
                                alternatives.append(Path(base_directory) / ntpath.basename(old_path))
                            alternatives.extend(Path(folder) / ntpath.basename(old_path) for folder in resources)
                            found = next((path for path in alternatives if path.is_file()), None)
                            resolved.append(str(found.resolve()) if found else old_path)
                            if found is None:
                                plan.missing_resources.append({'branch': sheet.title, 'id': row[0].value, 'path': old_path})
                        parameters['图像路径'] = '\n'.join(resolved)
                    row[2].value = json.dumps(parameters, ensure_ascii=False, allow_nan=False)
                repository.validate_workbook(converted)
                repository.import_from_workbook(converted)
                repository.execution_snapshot()
                plan.workbooks[sheet.title] = converted
                if any(row[1] == '发送消息' for row in converted['命令'].iter_rows(min_row=2, values_only=True)):
                    plan.warnings.append(f'分支“{sheet.title}”包含微信消息：已保留联系人及内容，执行需 Windows、受支持的微信客户端和 wxauto；迁移不发送消息。')
            except (ValueError, TypeError, WorkbookValidationError, GraphValidationError) as error:
                plan.errors.append({'sheet': sheet.title, 'error': str(error)})
            finally:
                split.close()
    if plan.default_branch not in plan.workbooks:
        plan.default_branch = next(iter(plan.workbooks), '')
    return plan


def _branch_filename(index, branch):
    safe = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', branch).strip(' .')[:60] or '分支'
    return f'{index:02d}-{safe}-新版.xlsx'


def write_migration_bundle(plan, source_path, output_parent=None):
    """Write a uniquely named migration folder; never replace user's files."""
    source = Path(source_path).resolve()
    parent = Path(output_parent).resolve() if output_parent else source.parent
    parent.mkdir(parents=True, exist_ok=True)
    folder = Path(tempfile.mkdtemp(prefix=f'{source.stem}-迁移-', dir=parent))
    shutil.copy2(source, folder / '原始备份.xlsx')
    outputs = {}
    original = load_workbook(source, read_only=True)
    try:
        branch_names = [sheet.title for sheet in original if sheet.title != '设置']
    finally:
        original.close()
    for branch, book in plan.workbooks.items():
        path = folder / _branch_filename(branch_names.index(branch)+1, branch)
        # Keep explicitly referenced images portable; never scan unrelated drives.
        assets = folder / 'images'
        rewritten = []
        for row in book['命令'].iter_rows(min_row=2):
            parameters = json.loads(row[2].value)
            if '图像路径' not in parameters:
                continue
            paths = []
            for value in str(parameters['图像路径']).splitlines():
                original_image = Path(value)
                if original_image.is_file():
                    assets.mkdir(exist_ok=True)
                    digest = hashlib.sha256(original_image.read_bytes()).hexdigest()[:16]
                    image_name = digest + original_image.suffix.lower()
                    destination = assets / image_name
                    if not destination.exists():
                        shutil.copy2(original_image, destination)
                    paths.append('images/' + image_name)
                else:
                    paths.append(value)
            parameters['图像路径'] = '\n'.join(paths)
            rewritten.append((row[2], row[2].value))
            row[2].value = json.dumps(parameters, ensure_ascii=False, allow_nan=False)
        try:
            book.save(path)
        finally:
            for cell, value in rewritten:
                cell.value = value
        outputs[branch] = str(path)
    report = {
        'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
        'source_commands': plan.source_commands, 'converted_commands': plan.converted_commands,
        'default_branch': plan.default_branch, 'outputs': outputs,
        'branches': plan.branches, 'errors': plan.errors, 'warnings': plan.warnings,
        'missing_resources': plan.missing_resources,
        'status': 'blocked' if plan.errors else 'converted_with_warnings' if plan.warnings or plan.missing_resources else 'converted',
    }
    (folder / '迁移报告.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    lines = ['An Clicker 旧版迁移报告', f'原指令：{plan.source_commands}；已转换：{plan.converted_commands}',
             f'原当前分支：{plan.default_branch}', '各分支独立；仅原有异常跳转会转到关联分支，不会顺序串联运行。',
             '请整体保留迁移目录：关联项目与 images 图片目录使用相对路径，不能只移动一个 Excel。',
             '使用：在 v1.1.5 或更新版本中导入对应的“新版.xlsx”。',
             '原始文件完整备份为“原始备份.xlsx”。', '', '生成文件：']
    lines.extend(f'{branch}：{Path(path).name}' for branch, path in outputs.items())
    lines += ['', '注意事项：', *plan.warnings]
    if plan.missing_resources:
        lines.append(f'有 {len(plan.missing_resources)} 条图片引用未找到文件。Excel 不包含图片本体，请复制旧图片目录或重新选择资源文件夹，未补齐前不要运行图像任务。')
    lines.extend(f'未转换：{item["sheet"]}：{item["error"]}' for item in plan.errors)
    (folder / '迁移说明.txt').write_text('\n'.join(lines), encoding='utf-8')
    return folder, outputs


def needs_batch_migration(source):
    if is_current_workbook(source):
        return False
    sheets = [sheet for sheet in source if sheet.title != '设置']
    return (len(sheets) > 1 or
            ('设置' in source.sheetnames and source['设置'].max_column == 2))


def choose_migration(window, source, source_path):
    """Interactive batch migration; return chosen file, never import or execute."""
    from qt_compat.QtWidgets import QInputDialog, QMessageBox
    from legacy_workbook import save_converted_copy

    plan = inspect_migration(source, Path(source_path).resolve().parent)
    try:
        summary = (f'共 {plan.source_commands} 条指令，已适配 {plan.converted_commands} 条，'
                   f'{len(plan.workbooks)} 个独立分支。\n'
                   f'缺少图片引用：{len(plan.missing_resources)} 条；转换错误：{len(plan.errors)} 项。\n'
                   '各分支另存独立新版文件，保留原始备份及迁移报告，不会串起来运行。\n'
                   '旧快捷键与窗口设置仅归档，不自动覆盖新版偏好。\n'
                   '是否生成迁移文件？')
        if QMessageBox.question(window, '旧版数据迁移器', summary,
                                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                                QMessageBox.StandardButton.No) != QMessageBox.StandardButton.Yes:
            return None
        folder, outputs = write_migration_bundle(plan, source_path)
        if plan.errors:
            QMessageBox.warning(window, '迁移未全部完成',
                                f'已生成可转换分支及报告，但有未适配内容，不自动导入。\n{folder}')
            return None
        names = list(outputs)
        if not names:
            QMessageBox.warning(window, '没有可导入指令', str(folder))
            return None
        selected, accepted = QInputDialog.getItem(
            window, '选择要导入的分支',
            '每个分支是独立任务。其他分支可随时通过“文件→导入”打开。',
            names, names.index(plan.default_branch), False,
        )
        if not accepted:
            QMessageBox.information(window, '迁移文件已保存', f'当前任务未改变。\n{folder}')
            return None
        backup = Workbook()
        try:
            window.workspace.repository.export_to_workbook(backup, window.db)
            save_converted_copy(backup, folder / '导入前任务备份.xlsx')
        finally:
            backup.close()
        QMessageBox.information(window, '迁移完成',
                                f'全部 {plan.converted_commands} 条指令已转换并保留在：\n{folder}\n'
                                f'即将导入：{selected}\n'
                                '缺失图片须另行补齐，迁移不会执行任何指令。')
        return outputs[selected]
    finally:
        plan.close()


def main():
    parser = argparse.ArgumentParser(description='离线迁移旧 Clicker Excel，不执行指令')
    parser.add_argument('source', type=Path)
    parser.add_argument('--output', type=Path)
    args = parser.parse_args()
    source = load_workbook(args.source)
    plan = None
    try:
        if is_current_workbook(source):
            raise ValueError('已经是新版工作簿，无需迁移')
        plan = inspect_migration(source, args.source.resolve().parent)
        folder, _ = write_migration_bundle(plan, args.source, args.output)
        print(folder)
        print(f'Converted {plan.converted_commands}/{plan.source_commands}; errors={len(plan.errors)}; missing images={len(plan.missing_resources)}')
        return 2 if plan.errors else 0
    finally:
        source.close()
        if plan is not None:
            plan.close()


if __name__ == '__main__':
    raise SystemExit(main())
