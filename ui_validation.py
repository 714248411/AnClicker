"""Release-only checks of actual loaded instruction dialogs, without input replay."""
from qt_compat.QtWidgets import QApplication, QAbstractSpinBox, QLineEdit
from instructions.registry import INSTRUCTION_SPECS
from instructions.models import ExecutionContext


def validate_instruction_editors(window):
    app = QApplication.instance()
    view = window.view_workspace
    original_theme = view.theme_mode
    checked = 0
    image_layout = False
    try:
        for theme in ('light', 'dark'):
            view.theme_mode = theme
            view._apply_theme()
            for spec in INSTRUCTION_SPECS:
                editor = spec.create_editor(window, context=ExecutionContext(metadata={'database': window.db}))
                try:
                    editor.show()
                    app.processEvents()
                    if spec.type_id == '图像点击':
                        image_layout = (hasattr(editor, 'folder_combo')
                                        and editor.ui.auxiliary_6.text() == '调整点击位置'
                                        and editor.capture_button.isVisible())
                        if not image_layout:
                            raise RuntimeError('Packaged image-click editor has obsolete layout')
                    for spin in editor.findChildren(QAbstractSpinBox):
                        if spin.isHidden():
                            continue
                        line = spin.findChild(QLineEdit)
                        if line is not None and line.height() < line.fontMetrics().height():
                            raise RuntimeError(f'{theme}/{spec.type_id}/{spin.objectName()}: '
                                               f'clipped digits: {line.height()} < {line.fontMetrics().height()}')
                        checked += 1
                finally:
                    editor.hide()
                    editor.deleteLater()
            app.processEvents()
    finally:
        view.theme_mode = original_theme
        view._apply_theme()
    return {'spinboxes_checked': checked, 'image_click_layout': image_layout,
            'image_click_modal_add': validate_image_click_modal_add(window)}


def validate_legacy_migration():
    """Exercise the bundled converter using only temporary files and databases."""
    import tempfile
    from pathlib import Path
    from openpyxl import Workbook, load_workbook
    from graph_repository import GraphRepository, WorkbookValidationError
    from legacy_workbook import LEGACY_HEADERS, convert_legacy_workbook, save_converted_copy
    from 数据库操作 import DatabaseOperation

    with tempfile.TemporaryDirectory(prefix='anclicker-migration-') as folder:
        path = str(Path(folder) / 'migration.db')
        DatabaseOperation(path)
        repository = GraphRepository(path)
        source = Workbook()
        source.active.title = '主流程'
        source.active.append(LEGACY_HEADERS)
        source.active.append([8, 'target.png', '图像点击',
                              "{'灰度': 'False', '精度': '0.8', '区域': '(0,0,0,0)'}",
                              None, None, None, 2, '提示异常并暂停', '图片'])
        source.active.append([2, '保留文字', '文本输入', "{'手动输入': 'False'}",
                              None, None, None, 1, '自动跳过', '文本'])
        source_path = Path(folder) / 'old.xlsx'
        source.save(source_path)
        original = source_path.read_bytes()
        converted = convert_legacy_workbook(source, folder)
        repository.validate_workbook(converted)
        output = save_converted_copy(converted, source_path)
        loaded = load_workbook(output)
        try:
            repository.import_from_workbook(loaded)
        finally:
            loaded.close()
        commands = repository.list_commands()
        if ([item.id for item in commands] != [8, 2]
                or commands[0].parameters['灰度'] is not False
                or commands[1].parameters['内容'] != '保留文字'
                or source_path.read_bytes() != original):
            raise RuntimeError('Packaged legacy conversion lost original values')
        before = repository.validate_graph()
        source.active.cell(2, 4).value = "__import__('os').system('invalid')"
        try:
            repository.import_from_workbook(source)
        except WorkbookValidationError:
            pass
        else:
            raise RuntimeError('Packaged importer accepted an executable expression')
        if repository.snapshot() != before:
            raise RuntimeError('Invalid legacy import changed existing data')
        source.close()
        converted.close()
    return {'commands_checked': 2, 'round_trip': True, 'rollback': True}


def validate_image_click_modal_add(window):
    """Exercise real modal confirmation and table insertion in the packaged app."""
    import tempfile
    from pathlib import Path
    from unittest.mock import patch
    from qt_compat.QtCore import QTimer, Qt
    from qt_compat.QtGui import QPixmap
    from qt_compat.QtWidgets import QDialog, QDialogButtonBox
    from instructions.registry import get_instruction_spec
    workspace = window.workspace
    spec = get_instruction_spec('图像点击')
    original_ids = {c.id for c in workspace.repository.list_commands()}
    with tempfile.TemporaryDirectory(prefix='anclicker-image-add-') as folder:
        image = Path(folder) / 'original.png'
        pixmap = QPixmap(40, 30)
        pixmap.fill(Qt.GlobalColor.red)
        if not pixmap.save(str(image)):
            raise RuntimeError('Cannot prepare image-add validation')
        editor = spec.create_editor(parent=window, context=workspace._editor_context())
        editor._set_image_path(image)
        class Picker(QDialog):
            def exec(self):
                QTimer.singleShot(250, lambda: editor.ui.buttonBox.button(QDialogButtonBox.StandardButton.Ok).click())
                return QDialog.DialogCode.Accepted
            def selected_region(self):
                return (0, 0, 40, 30)
            def save_selection(self, target):
                if not pixmap.save(str(target)):
                    raise RuntimeError('Cannot save validation capture')
        watchdog = QTimer()
        watchdog.setSingleShot(True)
        watchdog.timeout.connect(editor.reject)
        watchdog.start(5000)
        QTimer.singleShot(0, editor.capture_button.click)
        try:
            with patch.object(type(spec), 'create_editor', return_value=editor), \
                 patch('instructions.common.image_click_editor.SmartCaptureDialog', Picker):
                command_id = workspace.add_command('图像点击')
            if command_id is None:
                raise RuntimeError('Screenshot confirmation did not add command')
            window.view_workspace.refresh_table()
            table = window.view_workspace.command_table
            if not any(table.item(row, 0).data(Qt.ItemDataRole.UserRole) == command_id
                       for row in range(table.rowCount())):
                raise RuntimeError('Confirmed image command missing from table')
            command = workspace.repository.get_command(command_id)
            if command.parameters['图像路径'] == str(image):
                raise RuntimeError('Screenshot was not selected for new command')
        finally:
            watchdog.stop()
            created = [c.id for c in workspace.repository.list_commands() if c.id not in original_ids]
            if created:
                workspace.remove_commands(created, confirm=False)
            editor.deleteLater()
    return True
