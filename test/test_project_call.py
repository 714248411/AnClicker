"""Real workbook calls, isolated from desktop input and the user's database."""
import os
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
import threading
import time

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
import pytest
from openpyxl import Workbook
from PySide6.QtCore import QMimeData, QPoint, QPointF, Qt, QUrl
from PySide6.QtGui import QDragEnterEvent, QDropEvent
from PySide6.QtWidgets import QApplication
from graph_repository import GraphRepository
from instructions.models import ExecutionContext, InstructionDraft
from instructions.registry import get_instruction_spec
from main_work import CommandThread
from 数据库操作 import DatabaseOperation


def call(path, repeat=1):
    return InstructionDraft('运行项目', {'项目路径': str(path)}, repeat_count=repeat)


def mark(name):
    return InstructionDraft('文本输入', {'内容': name})


def project(path, drafts):
    db = DatabaseOperation(str(path.with_suffix('.db')))
    repo = GraphRepository(db.db_path)
    for draft in drafts:
        repo.add_command(draft, unconnected=True)
    book = Workbook()
    repo.export_to_workbook(book, db)
    book.save(path)
    book.close()
    return db, repo


def worker(db):
    with patch('main_work.DatabaseOperation', return_value=db):
        result = CommandThread(SimpleNamespace(execution_services={}))
    def error(command, exception):
        raise exception
    result._handle_command_error = error
    return result


def context(thread, path, log):
    result = ExecutionContext(services={'文本输入': lambda **kw: log.append(kw['command'].parameters['内容'])},
        metadata={'run_project': thread._run_project, 'project_path': str(path),
                  'project_stack': [path.resolve()], 'database': thread.db})
    thread._active_context = result
    return result


def test_nested_projects_return_without_expanding_or_mutating_root(tmp_path):
    c = tmp_path / '中文 C.xlsx'
    b = tmp_path / 'B.xlsx'
    a = tmp_path / 'A.xlsx'
    project(c, [mark('C')])
    project(b, [mark('B'), call(c.name)])
    db, repo = project(a, [mark('A'), call(b.name), mark('after')])
    before = repo.snapshot()
    original = {p: p.read_bytes() for p in (a, b, c)}
    thread = worker(db)
    log = []
    ctx = context(thread, a, log)
    metadata = ctx.metadata
    highlighted = []
    thread.send_type_and_id.connect(lambda *args: highlighted.append(args))
    thread._execute_commands(repo.list_commands(), ctx)
    assert log == ['A', 'B', 'C', 'after']
    assert repo.snapshot() == before
    assert len(repo.list_commands()) == 3
    assert len(highlighted) == 3
    assert ctx.metadata is metadata
    assert ctx.metadata['database'] is db
    assert all(p.read_bytes() == data for p, data in original.items())


def test_stop_child_prevents_remaining_child_and_parent_commands(tmp_path):
    b, a = tmp_path / 'B.xlsx', tmp_path / 'A.xlsx'
    project(b, [mark('stop'), mark('must not run')])
    db, repo = project(a, [call(b), mark('parent must not run')])
    thread = worker(db)
    log = []
    ctx = context(thread, a, log)
    def stop(**kw):
        log.append(kw['command'].parameters['内容'])
        thread.request_stop()
    ctx.services['文本输入'] = stop
    thread._execute_commands(repo.list_commands(), ctx)
    assert log == ['stop']
    assert ctx.stop_requested and not thread.start_state
    assert ctx.metadata['database'] is db


def test_recursive_call_blocked_and_metadata_restored(tmp_path):
    b, a = tmp_path / 'B.xlsx', tmp_path / 'A.xlsx'
    project(b, [call(a.name)])
    db, repo = project(a, [call(b.name)])
    thread = worker(db)
    ctx = context(thread, a, [])
    before = ctx.metadata
    with pytest.raises(ValueError, match='循环调用'):
        thread._execute_commands(repo.list_commands(), ctx)
    assert ctx.metadata is before


def test_repeated_project_calls_are_not_false_cycles(tmp_path):
    b, a = tmp_path / 'B.xlsx', tmp_path / 'A.xlsx'
    project(b, [mark('B')])
    db, repo = project(a, [call(b, repeat=3)])
    thread = worker(db)
    log = []
    thread._execute_commands(repo.list_commands(), context(thread, a, log))
    assert log == ['B'] * 3


@pytest.mark.parametrize('target', ['', 'absent.xlsx', 'wrong.txt'])
def test_invalid_path_does_not_change_root(tmp_path, target):
    a = tmp_path / 'A.xlsx'
    db, repo = project(a, [mark('A')])
    thread = worker(db)
    before = repo.snapshot()
    with pytest.raises(ValueError):
        thread._run_project(target, context(thread, a, []))
    assert repo.snapshot() == before


def test_stopped_call_never_loads_file(tmp_path):
    a = tmp_path / 'A.xlsx'
    db, _ = project(a, [])
    thread = worker(db)
    ctx = context(thread, a, [])
    thread.request_stop()
    with patch('openpyxl.load_workbook') as load:
        thread._run_project(str(tmp_path / 'missing.xlsx'), ctx)
    load.assert_not_called()


def test_editor_accepts_one_file_drop_and_roundtrips_one_command(tmp_path):
    app = QApplication.instance() or QApplication([])
    b = tmp_path / '中文 空格.xlsx'
    project(b, [mark('B')])
    editor = get_instruction_spec('运行项目').create_editor()
    try:
        control = editor._controls['项目路径']
        mime = QMimeData()
        mime.setUrls([QUrl.fromLocalFile(str(b))])
        enter = QDragEnterEvent(QPoint(3, 3), Qt.CopyAction, mime, Qt.LeftButton, Qt.NoModifier)
        QApplication.sendEvent(control, enter)
        assert enter.isAccepted()
        drop = QDropEvent(QPointF(3, 3), Qt.CopyAction, mime, Qt.LeftButton, Qt.NoModifier)
        QApplication.sendEvent(control, drop)
        assert drop.isAccepted()
        assert Path(editor.get_draft().parameters['项目路径']) == b
        assert not editor.test_button.isEnabled()
    finally:
        editor.close()
        editor.deleteLater()
        app.processEvents()


def test_real_worker_stops_during_child_wait_without_forced_termination(tmp_path):
    a, b = tmp_path / 'A.xlsx', tmp_path / 'B.xlsx'
    project(b, [mark('entered'), InstructionDraft('时间等待', {'时长': 60, '单位': '秒'}), mark('never')])
    db, _ = project(a, [call(b), mark('never parent')])
    thread = worker(db)
    entered = threading.Event()
    log = []
    def record(**kwargs):
        log.append(kwargs['command'].parameters['内容'])
        entered.set()
    thread.main_window.execution_services = {'文本输入': record}
    thread.start()
    try:
        assert entered.wait(5)
        time.sleep(0.1)
        started = time.monotonic()
        thread.request_stop()
        assert thread.wait(1500)
        assert time.monotonic() - started < 1.5
        assert log == ['entered']
    finally:
        thread.request_stop()
        thread.wait(5000)


def test_invalid_workbook_and_depth_limit_leave_context_intact(tmp_path):
    a, b = tmp_path / 'A.xlsx', tmp_path / 'invalid.xlsx'
    book = Workbook()
    book.save(b)
    book.close()
    db, repo = project(a, [mark('A')])
    thread = worker(db)
    ctx = context(thread, a, [])
    old = ctx.metadata
    before = repo.snapshot()
    with pytest.raises(ValueError):
        thread._run_project(str(b), ctx)
    assert ctx.metadata is old
    assert repo.snapshot() == before
    ctx.metadata['project_stack'] = [tmp_path / f'{i}.xlsx' for i in range(32)]
    with pytest.raises(ValueError, match='32'):
        thread._run_project(str(b), ctx)


def test_unsaved_root_requires_absolute_target(tmp_path):
    a = tmp_path / 'A.xlsx'
    db, _ = project(a, [])
    thread = worker(db)
    with pytest.raises(ValueError, match='完整路径'):
        thread._run_project('B.xlsx', ExecutionContext())
