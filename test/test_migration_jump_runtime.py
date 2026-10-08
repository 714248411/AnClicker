"""Anonymised old workbooks: transfers are neither flattened nor call-and-return."""
import json
from pathlib import Path
from unittest.mock import Mock, patch

import pytest
from openpyxl import Workbook, load_workbook

from legacy_workbook import LEGACY_HEADERS, convert_legacy_workbook, _convert_parameters
from migration_tool import inspect_migration, write_migration_bundle
from instructions.models import ExecutionContext, InstructionDraft, CommandRecord
from instructions.registry import get_instruction_spec
from test.test_flow_jumps import setup, app


def source_book(destination='主流程-3'):
    book = Workbook()
    sheet = book.active
    sheet.title = '主流程'
    sheet.append(LEGACY_HEADERS)
    for identifier, content, policy in [(17, 'A', destination), (21, 'B', '自动跳过'), (42, 'C', '自动跳过')]:
        sheet.append([identifier, content, '文本输入', 'None', None, None, None, 1, policy, None])
    child = book.create_sheet('分支-乙')
    child.append(LEGACY_HEADERS)
    child.append([5, 'child', '文本输入', '{}', None, None, None, 1, '自动跳过', None])
    return book


@pytest.mark.parametrize('destination,expected', [('主流程-3', ['A', 'C']), ('分支-乙-1', ['A', 'child'])])
def test_real_migrated_execution_preserves_jump_semantics(setup, tmp_path, destination, expected):
    db, repo, worker = setup
    source = source_book(destination)
    path = tmp_path / 'old.xlsx'
    source.save(path)
    plan = inspect_migration(source, tmp_path)
    assert not plan.errors and plan.converted_commands == 4
    folder, outputs = write_migration_bundle(plan, path)
    book = load_workbook(outputs['主流程'])
    repo.import_from_workbook(book)
    book.close()
    db.set_setting_value('当前文件路径', outputs['主流程'])
    events = []

    def execute(context, command):
        events.append(command.parameters['内容'])
        if events[-1] == 'A':
            raise ValueError('synthetic failure')

    context = ExecutionContext(services={'文本输入': execute})
    worker._active_context = context
    worker._execute_commands(repo.list_commands(), context)
    assert events == expected
    worker._handle_command_error.assert_not_called()
    # Same-branch destinations are ordinal rows, not ID 3 (IDs are 17/21/42).
    handler = repo.list_commands()[-1]
    assert handler.type_id == '报错跳转'
    editor = get_instruction_spec('报错跳转').create_editor(draft=handler.to_draft())
    assert editor.get_draft().parameters.get('旧版转移后结束') == handler.parameters.get('旧版转移后结束')
    editor.close()
    plan.close()


def test_invalid_jump_blocks_branch_and_retains_source(tmp_path):
    source = source_book('主流程-999')
    before = list(source['主流程'].values)
    plan = inspect_migration(source, tmp_path)
    assert plan.errors and '主流程' not in plan.workbooks
    assert before == list(source['主流程'].values)
    plan.close()


def test_portable_images_repeated_bundle_and_no_source_mutation(tmp_path):
    image = tmp_path / '中文图片.png'
    image.write_bytes(b'image fixture')
    source = Workbook()
    source.active.title = '主流程'
    source.active.append(LEGACY_HEADERS)
    source.active.append([8, str(image), '图像点击', '{}', None, None, None, 1, '自动跳过', None])
    old = tmp_path / 'old.xlsx'
    source.save(old)
    plan = inspect_migration(source, tmp_path)
    before = plan.workbooks['主流程']['命令'].cell(2, 3).value
    for _ in range(2):
        folder, paths = write_migration_bundle(plan, old)
        output = load_workbook(paths['主流程'])
        relative = json.loads(output['命令'].cell(2, 3).value)['图像路径']
        assert not Path(relative).is_absolute()
        assert (folder / relative).read_bytes() == image.read_bytes()
        output.close()
    assert plan.workbooks['主流程']['命令'].cell(2, 3).value == before
    plan.close()


def test_old_focus_and_unused_zero_count_parameters():
    parameters = _convert_parameters('窗口焦点等待', None,
        "{'标题包含':'Example','等待类型':'等待窗口失去焦点','检测频率':'1000','等待时间':60}", None)
    assert parameters['等待类型'] == '失去焦点' and parameters['检测频率'] == 1
    parameters = _convert_parameters('坐标点击', None,
        "{'动作':'右键单击','坐标':'470-170','自定义次数':0}", None)
    assert parameters['动作'] == '右键单击' and parameters['自定义次数'] == 1
    assert _convert_parameters('运行外部文件', 'C:/example.exe', 'None', None)['文件路径'] == 'C:/example.exe'


def test_message_data_preserved_without_native_sending(setup):
    db, repo, worker = setup
    parameters = _convert_parameters('发送消息', None, "{'联系人':'fixture','消息内容':'消息 ${x}'}", None)
    draft = InstructionDraft('发送消息', parameters)
    editor = get_instruction_spec('发送消息').create_editor(draft=draft)
    assert editor.get_draft().parameters == parameters
    editor.close()
    executor = get_instruction_spec('发送消息').create_executor()
    client = Mock()
    client.ChatWith.return_value = 'different person'
    with patch('sys.platform', 'win32'), patch.dict('sys.modules', {'wxauto': Mock(WeChat=lambda: client)}):
        with pytest.raises(RuntimeError, match='精确匹配'):
            executor.execute(ExecutionContext(), CommandRecord(1, '发送消息', parameters))
    client.SendMsg.assert_not_called()


def test_relative_image_preview_and_test_keep_portable_reference(setup, tmp_path):
    from PIL import Image
    from instructions.common import actions
    db, repo, worker = setup
    (tmp_path / 'images').mkdir()
    image = tmp_path / 'images' / '目标.png'
    Image.new('RGB', (24, 24), 'blue').save(image)
    project = tmp_path / 'task.xlsx'
    db.set_setting_value('当前文件路径', str(project))
    context = ExecutionContext(metadata={'database': db})
    editor = get_instruction_spec('图像点击').create_editor(
        draft=InstructionDraft('图像点击', {'图像路径': 'images/目标.png'}), context=context)
    assert editor._valid_image()
    assert editor.preview.pixmap() is not None and not editor.preview.pixmap().isNull()
    assert editor.get_draft().parameters['图像路径'] == 'images/目标.png'
    assert Path(actions.resolve_image_path(editor.get_draft().parameters, context)) == image
    editor.close()
