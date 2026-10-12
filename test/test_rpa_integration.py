import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import json
import threading
import time
from pathlib import Path
import pytest
from openpyxl import load_workbook
from instructions.common.rpa_bridge import RpaBridge
from instructions.models import ExecutionContext, InstructionDraft
from instructions.registry import INSTRUCTION_SPECS

@pytest.fixture
def bridge(tmp_path, monkeypatch):
    monkeypatch.setenv('ANCLICKER_DATA_DIR', str(tmp_path/'profile'))
    b=RpaBridge()
    yield b
    b.close()

def call(b, kind, params=None, variables=None):
    return b.execute({'id':kind,'params':params or {}}, variables)

def test_excel_roundtrip_copy_rows_saveas_and_plain_files(bridge, tmp_path):
    book=tmp_path/'源表.xlsx'
    call(bridge,'excelNew',{'path':str(book),'sheetName':'Sheet1'})
    call(bridge,'excelWrite',{'cell':'A1','value':'中文数据'})
    call(bridge,'excelWriteRange',{'range':'A2','data':'表格'}, {'表格':[[10,20],[30,40]]})
    assert call(bridge,'excelRead',{'cell':'A1','varName':'值'})['variables']['值']=='中文数据'
    assert call(bridge,'excelReadRange',{'range':'A2:B3','varName':'区域'})['variables']['区域']==[[10,20],[30,40]]
    call(bridge,'excelCopyPaste',{'from':'A2:B3','to':'D2','mode':'只粘贴值'})
    call(bridge,'excelInsertRow',{'row':1,'count':1})
    call(bridge,'excelDeleteRow',{'row':1,'count':1})
    call(bridge,'excelInsertCol',{'col':1,'count':1})
    call(bridge,'excelDeleteCol',{'col':1,'count':1})
    call(bridge,'excelSave',{'path':''})
    wb=load_workbook(book)
    assert wb.active['D3'].value==30
    assert wb.active['A1'].value=='中文数据'
    wb.close()
    call(bridge,'excelSaveAs',{'naming':'后缀递增','prefix':'','startIndex':1})
    assert (tmp_path/'源表-1.xlsx').is_file()
    destination=tmp_path/'目标.xlsx'
    call(bridge,'excelTableCopyPaste',{'fromPath':str(book),'fromRanges':'A1-A3','fromSheet':'Sheet1',
        'toPath':str(destination),'toSheet':'Sheet1','toCell':'B2','mode':'只粘贴值','save':'保存'})
    wb=load_workbook(destination)
    assert wb.active['B2'].value=='中文数据'
    assert wb.active['B4'].value==30
    wb.close()
    folder=tmp_path/'白话';folder.mkdir()
    (folder/'a.txt').write_text('alpha');(folder/'b.txt').write_text('beta')
    found=call(bridge,'fileFind',{'folder':str(folder),'keyword':'a','ext':'txt'})
    assert found['variables']['查找结果']==['a.txt']
    count=call(bridge,'fileCount',{'folder':str(folder),'ext':'txt'})
    assert count['variables']['统计结果']['文件个数']==2
    copy=tmp_path/'copy';copy.mkdir()
    call(bridge,'fileOp',{'action':'复制','folder':str(folder),'sortBy':'文件名','order':'从小到大','index':'1','dst':str(copy)})
    assert (copy/'a.txt').read_text()=='alpha'
    call(bridge,'fileRename',{'folder':str(copy),'sortBy':'文件名','order':'从小到大','index':'1','mode':'加前缀','text':'新_','type':'自定义文字','sep':'_','parts':''})
    assert (copy/'新_a.txt').is_file()

def test_stop_terminates_wait_and_engine_can_restart(bridge):
    context=ExecutionContext()
    timer=threading.Timer(.4,lambda:setattr(context,'stop_requested',True))
    # Start engine before measuring cancellation.
    call(bridge,'fileFind',{'folder':str(Path(__file__).parent),'keyword':'absent','ext':'txt'})
    timer.start();start=time.monotonic()
    with pytest.raises(InterruptedError):
        bridge.execute({'id':'browserWait','params':{'seconds':30}},context=context)
    timer.join()
    assert time.monotonic()-start<5
    assert bridge.process is None
    call(bridge,'fileFind',{'folder':str(Path(__file__).parent),'keyword':'absent','ext':'txt'})

def test_rpa_editors_roundtrip_reference_settings():
    from qt_compat.QtWidgets import QApplication
    app=QApplication.instance() or QApplication([])
    catalog=json.loads((Path(__file__).resolve().parents[1]/'integrations/yingdao/commands.json').read_text(encoding='utf-8'))
    by_id={s.type_id:s for s in INSTRUCTION_SPECS}
    for c in catalog:
        if c['type_id'] not in by_id:continue
        params={p['key']:p['default'] for p in c['params']}
        editor=by_id[c['type_id']].create_editor(draft=InstructionDraft(c['type_id'],params,3,'自动跳过','同一设置'))
        restored=editor.get_draft()
        assert restored.parameters==params
        assert (restored.repeat_count,restored.error_policy,restored.note)==(3,'自动跳过','同一设置')
        editor.close();editor.deleteLater()
    app.processEvents()

def test_browser_navigation_input_click_extract_screenshot_and_title(bridge, tmp_path):
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            body=b'''<!doctype html><title>Fixture title</title><input id="name"><button id="go" onclick="document.querySelector('#result').textContent=document.querySelector('#name').value">Run</button><div id="result">ready</div><p class="row">10</p><p class="row">20</p>'''
            self.send_response(200);self.send_header('Content-Type','text/html; charset=utf-8');self.end_headers();self.wfile.write(body)
        def log_message(self,*args):pass
    server=ThreadingHTTPServer(('127.0.0.1',0),Handler)
    threading.Thread(target=server.serve_forever,daemon=True).start()
    try:
        call(bridge,'openUrl',{'url':f'http://127.0.0.1:{server.server_port}','target':'新窗口','loginMode':'独立登录(每次全新)','engine':'跟随当前','waitType':'加载完成'})
        call(bridge,'waitElement',{'selector':'#name','by':'CSS选择器','timeout':3})
        call(bridge,'input',{'selector':'#name','by':'CSS选择器','text':'网页测试','clear':'是','pressEnter':'否'})
        call(bridge,'click',{'selector':'#go','by':'CSS选择器','clickType':'单击'})
        result=call(bridge,'getText',{'selector':'#result','by':'CSS选择器','varName':'输出'})
        assert result['variables']['输出']=='网页测试'
        result=call(bridge,'acInspect')
        assert result['result']['titles']==['极速网页AC']
        screenshot=tmp_path/'网页.png'
        call(bridge,'screenshot',{'selector':'#result','path':str(screenshot)})
        assert screenshot.read_bytes().startswith(b'\x89PNG')
        result=call(bridge,'scrapeText',{'selector':'.row','by':'CSS选择器','varName':'列表','all':'全部匹配'})
        assert result['variables']['列表']==['10','20']
        out=tmp_path/'网页结果.xlsx'
        call(bridge,'extractTextToTable',{'selector':'.row','by':'CSS选择器','all':'全部匹配','filter':'文本与数字','toPath':str(out),'toSheet':'Sheet1','toCell':'A1','save':'保存'})
        wb=load_workbook(out);assert str(wb.active['A1'].value)=='10';wb.close()
    finally:
        server.shutdown();server.server_close()

def test_quick_actions_replace_previous_children_and_preserve_parameters():
    from qt_compat.QtWidgets import QApplication
    from node_editor.palette import InstructionPalette
    app=QApplication.instance() or QApplication([])
    palette=InstructionPalette(INSTRUCTION_SPECS)
    calls=[]
    palette.quickCaptureRequested.connect(lambda:calls.append('capture'))
    palette.tree.itemClicked.emit(palette._quick_capture_item,0)
    assert calls==[]
    palette.tree.itemDoubleClicked.emit(palette._quick_capture_item,0)
    assert calls==['capture']
    palette.show_recent_capture('first.png')
    leaf=palette._quick_capture_item.child(0);palette.tree.setCurrentItem(leaf)
    assert palette.quick_parameters('图像点击')=={'图像路径':'first.png'}
    palette.show_recent_ocr((1,2,30,40))
    assert palette._quick_capture_item.childCount()==0
    assert palette._quick_ocr_item.childCount()==3
    palette.tree.setCurrentItem(palette._quick_ocr_item.child(1))
    assert palette.quick_parameters('OCR点击识别区域')=={'区域':'(1, 2, 30, 40)'}
    palette.search_edit.setText('复制识别文字')
    assert not palette._quick_ocr_item.isHidden()
    palette.show_recent_capture('second.png')
    assert palette._quick_ocr_item.childCount()==0
    assert palette._quick_capture_item.childCount()==1
    palette.close()


def test_native_executor_persistence_variables_and_picker_result(bridge, tmp_path, monkeypatch):
    from instructions.common import rpa_instruction
    from instructions.registry import get_instruction_spec
    from instructions.models import CommandRecord
    from graph_repository import GraphRepository
    from 数据库操作 import DatabaseOperation
    monkeypatch.setattr(rpa_instruction, 'get_bridge', lambda: bridge)
    db=DatabaseOperation(str(tmp_path/'host.db'))
    repo=GraphRepository(db.db_path)
    context=ExecutionContext(variables={'目标':str(tmp_path/'宿主.xlsx'),'内容':'宿主变量'})
    drafts=[InstructionDraft('新建Excel',{'path':'{{目标}}','sheetName':'Sheet1'}),
            InstructionDraft('Excel写入单元格',{'cell':'A1','value':'{{内容}}'}),
            InstructionDraft('读取单元格',{'cell':'A1','varName':'读回'}),
            InstructionDraft('保存Excel',{'path':''})]
    for draft in drafts: repo.add_command(draft)
    for command in GraphRepository(db.db_path).list_commands():
        get_instruction_spec(command.type_id).create_executor().execute(context,command)
    assert context.variables['读回']=='宿主变量'
    wb=load_workbook(tmp_path/'宿主.xlsx');assert wb.active['A1'].value=='宿主变量';wb.close()
    monkeypatch.setattr(bridge,'execute',lambda *a,**k:{'result':{'canceled':False,'css':'#target','xpath':'//input','text':''}})
    get_instruction_spec('拾取元素').create_executor().execute(context,CommandRecord(None,'拾取元素'))
    assert context.variables['网页元素']=='#target'


def test_file_rename_replacement_and_invalid_selection(bridge, tmp_path):
    file=tmp_path/'before.txt';file.write_text('retain')
    params={'folder':str(tmp_path),'sortBy':'文件名','order':'从小到大','index':'1',
            'mode':'替换文字','text':'before','replacement':'after','type':'自定义文字','sep':'_','parts':''}
    call(bridge,'fileRename',params)
    assert (tmp_path/'after.txt').read_text()=='retain'
    with pytest.raises(RuntimeError):
        call(bridge,'fileOp',{'action':'删除','folder':str(tmp_path),'index':'invalid','sortBy':'文件名','order':'从小到大'})
    assert (tmp_path/'after.txt').read_text()=='retain'


def test_open_url_waits_for_real_navigation_and_reports_failures(bridge):
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
    import socket
    agents=[]
    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            agents.append(self.headers.get('User-Agent',''))
            time.sleep(.7)
            body=b'<title>Loaded</title><div id="loaded">finished</div>'
            self.send_response(200);self.send_header('Content-Type','text/html; charset=utf-8');self.end_headers();self.wfile.write(body)
        def log_message(self,*args):pass
    server=ThreadingHTTPServer(('127.0.0.1',0),Handler)
    threading.Thread(target=server.serve_forever,daemon=True).start()
    try:
        for target in ('新窗口','当前标签页'):
            url=f'http://127.0.0.1:{server.server_port}/{target}'
            result=call(bridge,'openUrl',{'url':url,'target':target,'waitType':'加载完成'})
            assert result['result']['url'].startswith(f'http://127.0.0.1:{server.server_port}/')
            assert call(bridge,'getText',{'selector':'#loaded','by':'CSS选择器','varName':'value'})['variables']['value']=='finished'
        assert agents and all('Chrome/' in ua and 'Electron/' not in ua and 'anclicker-rpa-runtime/' not in ua for ua in agents)
        with socket.socket() as closed:
            closed.bind(('127.0.0.1',0));port=closed.getsockname()[1]
        with pytest.raises(RuntimeError,match='ERR_CONNECTION_REFUSED'):
            call(bridge,'openUrl',{'url':f'http://127.0.0.1:{port}','target':'新窗口'})
        with pytest.raises(RuntimeError,match='完整网页地址'):
            call(bridge,'openUrl',{'url':'https://','target':'新窗口'})
    finally:server.shutdown();server.server_close()


def test_clipboard_paste_to_web_and_excel(bridge,tmp_path):
    import pyperclip
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
    original=pyperclip.paste()
    text='识别后的中文文字 123'
    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            body=b'<input id="target"><button id="copy" onclick="document.querySelector(\'#result\').textContent=document.querySelector(\'#target\').value">Copy</button><div id="result"></div>'
            self.send_response(200);self.send_header('Content-Type','text/html; charset=utf-8');self.end_headers();self.wfile.write(body)
        def log_message(self,*args):pass
    server=ThreadingHTTPServer(('127.0.0.1',0),Handler)
    threading.Thread(target=server.serve_forever,daemon=True).start()
    try:
        pyperclip.copy(text)
        call(bridge,'openUrl',{'url':f'http://127.0.0.1:{server.server_port}','target':'新窗口'})
        call(bridge,'waitElement',{'selector':'#copy','by':'CSS选择器','timeout':2})
        call(bridge,'pasteToElement',{'selector':'#target','by':'CSS选择器','mode':'真实粘贴(Ctrl+V)'})
        call(bridge,'click',{'selector':'#copy','by':'CSS选择器'})
        assert call(bridge,'getText',{'selector':'#result','by':'CSS选择器','varName':'value'})['variables']['value']==text
        book=tmp_path/'粘贴结果.xlsx'
        call(bridge,'pasteToCell',{'toPath':str(book),'toSheet':'Sheet1','toCell':'B2','save':'保存'})
        wb=load_workbook(book);assert wb.active['B2'].value==text;wb.close()
    finally:
        pyperclip.copy(original);server.shutdown();server.server_close()
