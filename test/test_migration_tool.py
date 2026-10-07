import hashlib
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

import pytest
from openpyxl import Workbook, load_workbook

from graph_repository import GraphRepository
from legacy_workbook import LEGACY_HEADERS
from migration_tool import inspect_migration, write_migration_bundle, choose_migration, needs_batch_migration
from 数据库操作 import DatabaseOperation


def old_export():
    """Anonymised reproduction of the supplied two-branch / INI export."""
    book = Workbook()
    main = book.active
    main.title = '主流程'
    main.append(LEGACY_HEADERS + ['隶属分支'])
    main.append([174, None, '鼠标点击', "{'鼠标':'左键','次数':1,'间隔':100,'按压':1}",
                 None, None, None, 999, '提示异常并暂停', None, '主流程'])
    second = book.create_sheet('分支乙')
    second.append(LEGACY_HEADERS + ['隶属分支'])
    for index in range(40):
        second.append([134 + index, 'Z:/old/images/target.png', '图像点击',
                       "{'动作':'左键单击','异常':99,'区域':'(0,0,0,0)','灰度':False}",
                       None, None, None, 1, '提示异常并暂停', None, '分支乙'])
    settings = book.create_sheet('设置')
    for row in [('[Config]', None), ('图像匹配精度', '0.85'), ('当前分支', '分支乙'),
                ('时间间隔', '0.269'), (None, None), ('[分支]', None),
                ('主流程', "('W',1)"), ('分支乙', "('',3)"),
                ('[全局快捷键]', None), ('开始运行', 'f9'),
                ('[资源文件夹路径]', None), ('路径1', 'Z:/old/images')]:
        settings.append(row)
    return book


def test_complete_independent_branches_and_safe_settings(tmp_path):
    source = old_export()
    original = {sheet.title: list(sheet.values) for sheet in source}
    plan = inspect_migration(source, tmp_path)
    assert needs_batch_migration(source)
    assert not plan.errors
    assert plan.source_commands == plan.converted_commands == 41
    assert plan.default_branch == '分支乙'
    assert len(plan.workbooks['主流程']['命令']['A']) == 2
    assert plan.workbooks['主流程']['命令'].cell(2, 4).value == 999
    params = json.loads(plan.workbooks['分支乙']['命令'].cell(2, 3).value)
    assert params['灰度'] is False and params['异常'] == '99' and params['精度'] == .85
    assert len(plan.missing_resources) == 40
    settings = {row[1]: row[2] for row in plan.workbooks['分支乙']['设置'].iter_rows(min_row=2, values_only=True)}
    assert settings['运行重复次数'] == '3'
    assert '快捷键-开始运行' not in settings
    assert '时间间隔' in settings['旧版迁移配置归档']
    assert original == {sheet.title: list(sheet.values) for sheet in source}
    plan.close()


def test_bundle_original_and_all_outputs_validate(tmp_path):
    source = old_export()
    path = tmp_path / 'input.xlsx'
    source.save(path)
    digest = hashlib.sha256(path.read_bytes()).digest()
    plan = inspect_migration(source, tmp_path)
    folder, outputs = write_migration_bundle(plan, path)
    assert hashlib.sha256(path.read_bytes()).digest() == digest
    assert (folder / '原始备份.xlsx').read_bytes() == path.read_bytes()
    assert len(outputs) == 2
    db = DatabaseOperation(str(tmp_path / 'target.db'))
    repo = GraphRepository(db.db_path)
    for name, output in outputs.items():
        book = load_workbook(output)
        repo.import_from_workbook(book)
        snapshot = repo.validate_graph()
        assert len(snapshot.commands) == (1 if name == '主流程' else 40)
        book.close()
    report = json.loads((folder / '迁移报告.json').read_text(encoding='utf-8'))
    assert report['converted_commands'] == 41
    assert report['status'] == 'converted_with_warnings'
    plan.close()


def test_resource_lookup_does_not_execute_task(tmp_path):
    (tmp_path / 'target.png').write_bytes(b'fixture')
    with patch('os.system', side_effect=AssertionError('must not execute')):
        plan = inspect_migration(old_export(), tmp_path)
    assert not plan.missing_resources
    params = json.loads(plan.workbooks['分支乙']['命令'].cell(2, 3).value)
    assert params['图像路径'] == str(tmp_path / 'target.png')
    plan.close()


def test_errors_do_not_hide_other_branches(tmp_path):
    book = old_export()
    book['主流程'].cell(2, 3).value = '未知指令'
    plan = inspect_migration(book, tmp_path)
    assert plan.source_commands == 41 and plan.converted_commands == 40
    assert len(plan.errors) == 1
    assert plan.errors[0]['sheet'] == '主流程'
    assert '分支乙' in plan.workbooks
    plan.close()


def test_cross_branch_jump_not_flattened(tmp_path):
    book = old_export()
    book['分支乙'].cell(2, 9).value = '主流程-1'
    plan = inspect_migration(book, tmp_path)
    assert '分支乙' not in plan.workbooks
    assert plan.errors
    plan.close()


def test_ui_selection_keeps_current_database_until_import(tmp_path):
    from PySide6.QtWidgets import QInputDialog, QMessageBox
    source = old_export()
    path = tmp_path / 'old.xlsx'
    source.save(path)
    database = DatabaseOperation(str(tmp_path / 'current.db'))
    repo = GraphRepository(database.db_path)
    before = repo.snapshot()
    window = SimpleNamespace(db=database, workspace=SimpleNamespace(repository=repo))
    with patch.object(QMessageBox, 'question', return_value=QMessageBox.StandardButton.Yes), \
         patch.object(QMessageBox, 'information'), \
         patch.object(QInputDialog, 'getItem', return_value=('分支乙', True)) as chooser:
        result = choose_migration(window, source, str(path))
    assert Path(result).is_file()
    assert repo.snapshot() == before
    assert chooser.call_args.args[4] == 1  # Old current branch selected by default.
    assert list(Path(result).parent.glob('导入前任务备份-新版*.xlsx'))


def test_ui_cancel_produces_no_files(tmp_path):
    from PySide6.QtWidgets import QMessageBox
    source = old_export()
    path = tmp_path / 'old.xlsx'
    source.save(path)
    with patch.object(QMessageBox, 'question', return_value=QMessageBox.StandardButton.No):
        assert choose_migration(None, source, str(path)) is None
    assert list(tmp_path.iterdir()) == [path]
