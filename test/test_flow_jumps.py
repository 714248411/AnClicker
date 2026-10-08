import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
from types import SimpleNamespace
from unittest.mock import patch, Mock
import pytest
from openpyxl import Workbook, load_workbook
from PySide6.QtWidgets import QApplication, QWidget
from PySide6.QtCore import QPointF
from graph_repository import GraphRepository, GraphValidationError
from instructions.models import InstructionDraft, ExecutionContext
from instructions.registry import get_instruction_spec, INSTRUCTION_SPECS
from main_work import CommandThread
from 数据库操作 import DatabaseOperation
from flow_jumps import compare_variables, target_config
from node_editor import NodeEditorWidget


@pytest.fixture
def app():
    app=QApplication.instance() or QApplication([])
    yield app
    app.processEvents()

@pytest.fixture
def setup(tmp_path,app):
    db=DatabaseOperation(str(tmp_path/'test.db'))
    repo=GraphRepository(db.db_path)
    with patch('main_work.DatabaseOperation',return_value=db):
        worker=CommandThread(SimpleNamespace(execution_services={}))
    worker._persist_variables=Mock()
    worker._handle_command_error=Mock(return_value='continue')
    return db,repo,worker

def add(repo,name='文本输入',**parameters):
    command=repo.add_command(InstructionDraft(name,parameters),unconnected=True)
    node=next(n.node_id for n in repo.snapshot().nodes if n.command_id==command.id)
    return node

def run(repo,worker,log,fail=(),variables=None):
    def text(**kw):
        name=kw['command'].parameters['内容']
        log.append(name)
        if name in fail: raise ValueError('test failure')
    context=ExecutionContext(variables=variables or {},services={'文本输入':text},output=lambda s:None)
    worker._active_context=context
    worker._execute_commands(repo.list_commands(),context)
    return context

def fixture_error(repo,mode='连线节点',row=3,path=''):
    a=add(repo,内容='A')
    b=add(repo,内容='B')
    c=add(repo,内容='C')
    handler=add(repo,'报错跳转',跳转方式=mode,目标行=row,项目路径=path)
    repo.connect_nodes(a,handler)
    if mode=='连线节点': repo.connect_nodes(handler,c)
    return a,b,c,handler

@pytest.mark.parametrize('mode',['连线节点','当前项目行'])
def test_error_jump_skips_old_path_and_starts_target_row(setup,mode):
    db,repo,worker=setup
    fixture_error(repo,mode)
    log=[]
    run(repo,worker,log,fail={'A'})
    assert log==['A','C']
    worker._handle_command_error.assert_not_called()

def test_no_error_preserves_default_order_and_does_not_execute_handler(setup):
    db,repo,worker=setup
    fixture_error(repo)
    log=[]
    run(repo,worker,log)
    assert log==['A','B','C']

def test_disconnected_handler_never_enables_jump(setup):
    db,repo,worker=setup
    a,b,c,handler=fixture_error(repo)
    repo.delete_node_connections(a,'outgoing')
    log=[]
    run(repo,worker,log,fail={'A'})
    assert log==['A','B','C']
    worker._handle_command_error.assert_called_once()

@pytest.mark.parametrize('mode,row',[('当前项目行',999),('连线节点',1)])
def test_invalid_target_stops_without_running_old_path(setup,mode,row):
    db,repo,worker=setup
    a,b,c,handler=fixture_error(repo,mode,row)
    repo.delete_node_connections(handler,'outgoing')
    log=[]
    run(repo,worker,log,fail={'A'})
    assert log==['A']
    assert not worker.start_state

def test_backward_error_jump_can_retry_and_is_interruptible(setup):
    db,repo,worker=setup
    a=add(repo,内容='A')
    handler=add(repo,'报错跳转',跳转方式='连线节点',目标行=1)
    repo.connect_nodes(a,handler)
    repo.connect_nodes(handler,a)
    repo.execution_snapshot() # Backward exceptional edges are not normal DAG cycles.
    calls=[]
    def action(**kw):
        calls.append('A')
        if len(calls)==3: worker.request_stop()
        raise ValueError('retry')
    context=ExecutionContext(services={'文本输入':action})
    worker._active_context=context
    worker._execute_commands(repo.list_commands(),context)
    assert calls==['A']*3 and context.stop_requested

def write_project(path,names):
    db=DatabaseOperation(str(path.with_suffix('.db')))
    repo=GraphRepository(db.db_path)
    for name in names: add(repo,内容=name)
    book=Workbook()
    repo.export_to_workbook(book,db)
    book.save(path)
    book.close()

def test_cross_project_from_row_and_return_without_mutating_workspace(setup,tmp_path):
    db,repo,worker=setup
    child=tmp_path/'子项目.xlsx'
    write_project(child,['skip child','child 2','child 3'])
    fixture_error(repo,'其他项目行',2,str(child))
    before=repo.snapshot()
    log=[]
    run(repo,worker,log,fail={'A'})
    assert log==['A','child 2','child 3','B','C']
    assert repo.snapshot()==before

@pytest.mark.parametrize('mode',['当前项目行','连线节点'])
@pytest.mark.parametrize('value,expected',[(10,['yes','after']),(0,['no','yes','after'])])
def test_variable_comparison_routes_selected_branch(setup,mode,value,expected):
    db,repo,worker=setup
    node=add(repo,'变量比较',变量='count',比较条件='大于',比较对象='固定值',比较值='5',比较类型='数字',
             是跳转方式=mode,是目标行=3,否跳转方式=mode,否目标行=2)
    no=add(repo,内容='no')
    yes=add(repo,内容='yes')
    after=add(repo,内容='after')
    if mode=='连线节点':
        repo.connect_nodes(node,yes,7)
        repo.connect_nodes(node,no,8)
    log=[]
    run(repo,worker,log,variables={'count':value})
    assert log==expected

@pytest.mark.parametrize('operator,right,expected',[('等于','10',True),('不等于','10',False),('大于','9',True),
                         ('小于','11',True),('大于等于','10',True),('小于等于','10',True),('包含','0',True),('不包含','2',True)])
def test_comparison_operators_and_right_variable(operator,right,expected):
    assert compare_variables({'变量':'a','比较条件':operator,'比较值':'b','比较对象':'变量'}, {'a':10,'b':right}) is expected

def test_comparison_invalid_variables_and_numbers():
    with pytest.raises(ValueError): compare_variables({'变量':'missing'}, {})
    with pytest.raises(ValueError): compare_variables({'变量':'a','比较值':'nan','比较类型':'数字'},{'a':1})
    with pytest.raises(ValueError): target_config({'跳转方式':'其他项目行','项目路径':''})
    with pytest.raises(ValueError): target_config({'目标行':1.2})

def test_edge_kinds_roundtrip_and_validation(setup,tmp_path):
    db,repo,worker=setup
    a,b,c,handler=fixture_error(repo)
    assert {edge.kind for edge in repo.snapshot().edges if edge.kind>=5}=={5,6}
    with pytest.raises(GraphValidationError): repo.connect_nodes(b,c,5)
    with pytest.raises(GraphValidationError): repo.connect_nodes(a,b,7)
    with pytest.raises(GraphValidationError): repo.connect_nodes(handler,b) # second destination is ambiguous
    book=Workbook()
    repo.export_to_workbook(book,db)
    restored=GraphRepository(DatabaseOperation(str(tmp_path/'restored.db')).db_path)
    restored.import_from_workbook(book)
    assert restored.snapshot().edges==repo.snapshot().edges
    restored.execution_snapshot()

def test_default_error_ports_and_dash_scale_independence(setup,app):
    db,repo,worker=setup
    fixture_error(repo)
    editor=NodeEditorWidget()
    editor.load_graph(repo.snapshot().nodes,repo.snapshot().edges,INSTRUCTION_SPECS,allow_incomplete=True)
    for node in editor.scene.nodes_by_id.values():
        assert (node.error_port is not None)==(not node.is_terminal and node.type_id!='报错跳转')
    edges=[e for e in editor.scene.edges if e.link_kind>=5]
    viewport=QWidget()
    viewport.resize(800,600)
    before=edges[0].jump_pen(viewport)
    editor.view._set_zoom(2)
    after=edges[0].jump_pen(viewport)
    assert before.isCosmetic() and before.dashPattern()==after.dashPattern()
    viewport.resize(1200,900)
    assert edges[0].jump_pen(viewport).dashPattern()!=before.dashPattern()
    for edge in edges:
        assert edge.path().boundingRect().top() < min(n.sceneBoundingRect().top() for n in editor.scene.nodes_by_id.values()) or edge.path().boundingRect().bottom() > max(n.sceneBoundingRect().bottom() for n in editor.scene.nodes_by_id.values())
    editor.close()
    viewport.close()

def test_jump_editors_roundtrip_modes_and_variable_dropdown(setup):
    db,repo,worker=setup
    context=ExecutionContext(variables={'score':10},metadata={'database':db})
    for type_id in ('报错跳转','变量比较'):
        editor=get_instruction_spec(type_id).create_editor(context=context)
        assert not editor.test_button.isEnabled()
        for prefix in editor.PREFIXES:
            editor._controls[prefix+'跳转方式'].setCurrentText('其他项目行')
            editor._controls[prefix+'项目路径'].setText('relative.xlsx')
            editor._controls[prefix+'目标行'].setValue(3)
        if type_id=='变量比较': editor._controls['变量'].setCurrentText('score')
        draft=editor.get_draft()
        editor.load_draft(draft)
        assert editor.get_draft()==draft
        editor.close()

@pytest.mark.parametrize('diamond',[False,True])
@pytest.mark.parametrize('cross_body',[False,True])
def test_arrow_anchor_stays_outside_target_and_follows_incoming_direction(app,diamond,cross_body):
    import math
    from PySide6.QtGui import QColor,QPainterPath
    from PySide6.QtWidgets import QGraphicsScene
    from node_editor.items import NodeItem,EdgeItem
    scene=QGraphicsScene()
    source=NodeItem('a',1,'文本输入','A',QColor('#336699'))
    target=NodeItem('b',2,'条件判断' if diamond else '文本输入','B',QColor('#336699'),control_kind_='condition' if diamond else None)
    scene.addItem(source)
    scene.addItem(target)
    target.setPos(500,200)
    edge=EdgeItem(source,target)
    scene.addItem(edge)
    end=target.input_port.scenePos()
    x=target.scenePos().x()+target.width+200 if cross_body else end.x()-200
    path=QPainterPath(QPointF(x,end.y()))
    path.lineTo(end)
    edge.setPath(path)
    tip,angle=edge.arrow_anchor()
    body=target.mapToScene(target.shape())
    assert not body.contains(tip)
    if cross_body:
        assert tip.x()>target.scenePos().x()+target.width
        assert math.cos(angle)<-0.99
    else:
        assert tip.x()<end.x()
        assert math.cos(angle)>0.99

@pytest.mark.parametrize('diamond',[False,True])
@pytest.mark.parametrize('direction',[(0,-1),(0,1),(-1,-1),(1,1)])
def test_arrow_top_bottom_diagonal_entries(app,diamond,direction):
    import math
    from PySide6.QtGui import QColor,QPainterPath
    from PySide6.QtWidgets import QGraphicsScene
    from node_editor.items import NodeItem,EdgeItem
    scene=QGraphicsScene()
    source=NodeItem('a',1,'文本输入','A',QColor('#336699'))
    target=NodeItem('b',2,'文本输入','B',QColor('#336699'),control_kind_='condition' if diamond else None)
    scene.addItem(source)
    scene.addItem(target)
    target.setPos(500,200)
    edge=EdgeItem(source,target)
    scene.addItem(edge)
    center=target.mapToScene(target.shape().boundingRect().center())
    dx,dy=direction
    path=QPainterPath(center+QPointF(dx*400,dy*400))
    path.lineTo(center)
    edge.setPath(path)
    tip,angle=edge.arrow_anchor()
    body=target.mapToScene(target.shape())
    assert not body.contains(tip)
    assert body.contains(tip+QPointF(math.cos(angle)*5,math.sin(angle)*5))
    assert math.cos(angle)*dx+math.sin(angle)*dy < -.99
