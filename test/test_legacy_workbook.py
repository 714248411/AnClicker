import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

import pytest
from openpyxl import Workbook, load_workbook

from graph_repository import GraphRepository, WorkbookValidationError
from legacy_workbook import (
    LEGACY_HEADERS, DATABASE_HEADERS, convert_legacy_workbook, save_converted_copy,
)
from 数据库操作 import DatabaseOperation


def legacy(rows, headers=LEGACY_HEADERS):
    book = Workbook()
    book.active.title = "主流程"
    book.active.append(headers)
    for row in rows:
        book.active.append(row)
    return book


def row(type_id="图像点击", payload="目标.png", parameters=None, id=5):
    return [id, payload, type_id, repr(parameters or {"灰度": "False", "精度": "0.8", "区域": "(0,0,0,0)"}),
            None, None, None, 2, "提示异常并暂停", "原始备注"]


@pytest.fixture
def repository(tmp_path):
    path = str(tmp_path / "test.db")
    DatabaseOperation(path)
    return GraphRepository(path)


@pytest.mark.parametrize("headers", [LEGACY_HEADERS, DATABASE_HEADERS])
def test_convert_import_export_and_order(repository, headers):
    source = legacy([row(id=50), row("文本输入", "原始文字☾变量☽", {"手动输入": "False"}, id=3)], headers)
    original = list(source.active.values)
    repository.import_from_workbook(source)
    commands = repository.list_commands()
    assert [item.id for item in commands] == [50, 3]
    assert [item.order for item in commands] == [0, 1]
    assert commands[0].parameters["灰度"] is False
    assert commands[0].parameters["区域"] == ""
    assert commands[1].parameters == {"内容": "原始文字☾变量☽", "手动输入": False}
    assert commands[0].repeat_count == 2
    assert commands[0].note == "原始备注"
    snapshot = repository.validate_graph()
    assert len(snapshot.nodes) == 4
    assert len(snapshot.edges) == 3
    assert list(source.active.values) == original
    exported = Workbook()
    repository.export_to_workbook(exported)
    assert convert_legacy_workbook(exported) is exported
    repository.import_from_workbook(exported)
    assert repository.list_commands() == commands


@pytest.mark.parametrize("type_id,payload,parameters,key,expected", [
    ("多图点击", "a.png、b.png", {"灰度": "False"}, "图像路径", "a.png\nb.png"),
    ("按下键盘", None, {"按键": "ctrl+c", "按压时长": "50"}, "按压时长", 50),
    ("时间等待", None, {"类型": "随机等待", "最小": "10-毫秒", "最大": "2-秒"}, "最小单位", "毫秒"),
    ("鼠标拖拽", None, {"开始位置": "1,2", "结束位置": "3,4", "移动速度": 2500, "开始随机": "False"}, "移动速度", 2.5),
    ("窗口焦点等待", None, {"标题包含": "测试", "检测频率": 500}, "检测频率", 0.5),
    ("滚轮滑动", None, {"方向": "↑"}, "方向", "向上"),
    ("获取Excel", "book.xlsx", {"递增": "False", "变量": "值"}, "工作簿", "book.xlsx"),
    ("运行Python", "raise RuntimeError('不要执行')", {"变量": "结果"}, "代码", "raise RuntimeError('不要执行')"),
])
def test_parameter_conversion(type_id, payload, parameters, key, expected):
    converted = convert_legacy_workbook(legacy([row(type_id, payload, parameters)]))
    assert json.loads(converted["命令"].cell(2, 3).value)[key] == expected


def test_relative_resource_and_no_overwrite(tmp_path):
    image = tmp_path / "目标.png"
    image.touch()
    original = tmp_path / "旧.xlsx"
    source = legacy([row()])
    source.save(original)
    before = original.read_bytes()
    converted = convert_legacy_workbook(source, tmp_path)
    assert json.loads(converted["命令"].cell(2, 3).value)["图像路径"] == str(image)
    first = save_converted_copy(converted, original)
    second = save_converted_copy(converted, original)
    assert first != second
    assert original.read_bytes() == before
    with Path(first).open("rb") as stream:
        book = load_workbook(stream)
        assert book.sheetnames == ["命令", "节点", "连线", "设置"]
        book.close()


@pytest.mark.parametrize("bad", [
    [5, "a.png", "图像点击", "__import__('os').system('echo unsafe')", None, None, None, 1, "自动跳过", ""],
    [5, "a.png", "未知指令", "{}", None, None, None, 1, "自动跳过", ""],
    [5, "a.png", "图像点击", "{}", None, None, None, 1, "分支2-1", ""],
    [5, "a.png", "图像点击", "{}", "unsupported", None, None, 1, "自动跳过", ""],
    [5, "a.png", "图像点击", "{'灰度':'maybe'}", None, None, None, 1, "自动跳过", ""],
    [5, "a.png", "图像点击", "{'精度':999}", None, None, None, 1, "自动跳过", ""],
    [5, "a.png", "图像点击", "{'未适配参数':10}", None, None, None, 1, "自动跳过", ""],
    [5, "", "图像点击", "{}", None, None, None, 1, "自动跳过", ""],
])
def test_reject_atomically_with_row_number(repository, bad):
    repository.import_from_workbook(legacy([row()]))
    before = repository.snapshot()
    with patch("os.system", side_effect=AssertionError("must not execute")):
        with pytest.raises(WorkbookValidationError, match="第 2 行"):
            repository.import_from_workbook(legacy([bad]))
    assert repository.snapshot() == before


def test_multiple_branches_are_not_flattened(repository):
    source = legacy([row()])
    other = source.create_sheet("分支2")
    other.append(LEGACY_HEADERS)
    other.append(row(id=9))
    with pytest.raises(WorkbookValidationError, match="多个非空分支"):
        repository.import_from_workbook(source)
    assert not repository.list_commands()


def test_legacy_branch_sheet_and_settings(repository):
    source = legacy([row() + ["主流程"]], LEGACY_HEADERS + ["隶属分支"])
    settings = source.create_sheet("设置")
    settings.append(["类型", "名称", "值", "附加值", "排序"])
    settings.append(["分支", "主流程", "", 1, 0])
    settings.append(["资源文件夹", "C:/images", None, None, 0])
    repository.import_from_workbook(source)
    assert len(repository.list_commands()) == 1
    settings.cell(2, 4).value = 10
    with pytest.raises(WorkbookValidationError, match="整组重复"):
        repository.import_from_workbook(source)


def test_invalid_new_graph_does_not_fall_back(repository):
    source = Workbook()
    repository.export_to_workbook(source)
    source["连线"].cell(2, 2).value = "missing"
    with pytest.raises(WorkbookValidationError):
        repository.import_from_workbook(source)


@pytest.mark.parametrize("accept", [True, False])
def test_import_dialog_converts_and_keeps_original(tmp_path, accept):
    from Start_Win import Main_window, QMessageBox

    database = DatabaseOperation(str(tmp_path / "ui.db"))
    repo = GraphRepository(database.db_path)
    repo.import_from_workbook(legacy([row(id=99)]))
    source = tmp_path / "旧指令.xlsx"
    legacy([row(id=2)]).save(source)
    before = source.read_bytes()
    window = SimpleNamespace(
        db=database, workspace=SimpleNamespace(repository=repo, reload_graph=Mock(), clear_connection_history=Mock()),
        view_workspace=SimpleNamespace(refresh_all=Mock()), menuzv=Mock(),
        add_recent_to_fileMenu=Mock(), statusBar=Mock(),
    )
    answer = QMessageBox.StandardButton.Yes if accept else QMessageBox.StandardButton.No
    with patch.object(QMessageBox, "question", return_value=answer), \
         patch.object(QMessageBox, "information"), patch.object(QMessageBox, "warning") as warning:
        Main_window.data_import(window, str(source))
        assert window.workspace.clear_connection_history.call_count == int(accept)
    warning.assert_not_called()
    assert source.read_bytes() == before
    assert [item.id for item in repo.list_commands()] == ([2] if accept else [99])
    if accept:
        saved = Path(database.get_setting_value("当前文件路径"))
        assert saved != source and saved.is_file()
        backups = list(tmp_path.glob("导入前指令备份-新版*.xlsx"))
        assert len(backups) == 1
        backup = load_workbook(backups[0])
        assert backup["命令"].cell(2, 1).value == 99
        backup.close()
        window.view_workspace.refresh_all.assert_called_once()
    else:
        assert list(tmp_path.glob("*-新版*.xlsx")) == []
        window.view_workspace.refresh_all.assert_not_called()
