import json
from openpyxl import Workbook, load_workbook
from qt_compat.QtWidgets import QApplication
from test.test_hidden_stop import host
from instructions.models import InstructionDraft
from instructions.registry import get_instruction_spec
from graph_repository import GraphRepository
from 数据库操作 import DatabaseOperation
from background_catalog import label


def test_background_commands_appear_in_table_graph_and_workbook(host,tmp_path):
    repo=host.workspace.repository;repo.clear()
    drafts=[InstructionDraft('绑定选择窗口',{'窗口句柄':'123','窗口标题':'测试窗口'}),
            InstructionDraft('后台图色操作',{'接口':label('GetColor'),'参数':'{"x":10,"y":20}','变量':'颜色'}),
            InstructionDraft('文本输入',{'内容':'后台中文'}),InstructionDraft('解除后台绑定',{})]
    commands=[repo.add_command(draft,unconnected=True) for draft in drafts]
    nodes={n.command_id:n.node_id for n in repo.snapshot().nodes if n.command_id is not None}
    repo.connect_nodes('start',nodes[commands[0].id])
    for left,right in zip(commands,commands[1:]):repo.connect_nodes(nodes[left.id],nodes[right.id])
    repo.connect_nodes(nodes[commands[-1].id],'end')
    host.workspace.reload_graph();host.view_workspace.refresh_table();QApplication.processEvents()
    table=host.view_workspace.command_table
    assert table.rowCount()==len(drafts)
    assert [table.item(i,2).text() for i in range(table.rowCount())]==[d.type_id for d in drafts]
    assert 'GetColor' in table.item(1,3).text()
    assert all(host.workspace.focus_command(c.id) for c in commands)
    repo.validate_graph()
    book=Workbook();repo.export_to_workbook(book,host.db)
    file=tmp_path/'后台项目.xlsx';book.save(file);book.close()
    imported=GraphRepository(DatabaseOperation(str(tmp_path/'imported.db')).db_path)
    book=load_workbook(file)
    try:imported.import_from_workbook(book)
    finally:book.close()
    assert [(c.type_id,c.parameters) for c in imported.list_commands()]==[(d.type_id,d.parameters) for d in drafts]
    assert [(e.source,e.target) for e in imported.snapshot().edges]==[(e.source,e.target) for e in repo.snapshot().edges]
    for draft in drafts:
        editor=get_instruction_spec(draft.type_id).create_editor(draft=draft)
        assert editor.get_draft().type_id==draft.type_id
        editor.deleteLater()


def test_background_component_settings_opens_with_real_controls(host):
    from background_component_settings import ComponentSettings
    from qt_compat.QtWidgets import QLineEdit,QComboBox
    editor=get_instruction_spec('绑定选择窗口').create_editor(context=host.workspace._editor_context())
    dialog=ComponentSettings(editor)
    assert dialog.findChildren(QLineEdit)
    assert dialog.findChild(QComboBox).currentText()=='7.2424 独立组件'
    dialog.deleteLater();editor.deleteLater()
