"""Convert recognised Clicker Excel exports without executing legacy expressions.

The source workbook is never modified. Branch jumps need an explicit graph
translation; they must not be flattened into a sequential automation.
"""
from __future__ import annotations

import ast
import json
import math
import re
from pathlib import Path

from openpyxl import Workbook

from graph_repository import (
    COMMAND_SHEET_HEADERS, NODE_SHEET_HEADERS, EDGE_SHEET_HEADERS,
    SETTINGS_SHEET_HEADERS, WORKBOOK_SHEETS, START_NODE_ID, END_NODE_ID,
    WorkbookValidationError,
)
from instructions.registry import get_instruction_spec

LEGACY_HEADERS = [
    "ID", "图像名称", "指令类型", "参数信息", "参数-2", "参数-3",
    "参数-4", "重复次数", "异常处理", "备注",
]
DATABASE_HEADERS = [
    "ID", "图像名称", "指令类型", "参数1", "参数2", "参数3",
    "参数4", "重复次数", "异常处理", "备注",
]
PAYLOAD_FIELDS = {
    "图像点击": "图像路径", "多图点击": "图像路径", "图像等待": "图像路径",
    "信息录入": "图像路径", "文本输入": "内容", "获取Excel": "工作簿",
    "写入单元格": "工作簿", "屏幕截图": "保存路径", "运行Python": "代码",
    "运行cmd": "命令", "运行外部文件": "文件路径",
}
SUPPORTED_TYPES = set(PAYLOAD_FIELDS) | {
    "坐标点击", "移动鼠标", "鼠标点击", "滚轮滑动", "按下键盘",
    "鼠标拖拽", "时间等待", "按键等待", "窗口焦点等待",
    "获取时间", "获取鼠标位置", "获取剪切板", "获取对话框",
    "数字验证码", "OCR识别", "窗口控制",
    "发送消息",
}
POLICIES = {"自动跳过", "提示异常并暂停", "提示异常并停止"}


def is_current_workbook(workbook):
    return set(workbook.sheetnames) == set(WORKBOOK_SHEETS) and list(
        next(workbook["命令"].iter_rows(max_row=1, values_only=True))
    ) == COMMAND_SHEET_HEADERS


def _integer(value, label, minimum=1):
    if isinstance(value, bool):
        raise ValueError(f"{label}必须是整数")
    number = float(value)
    if not math.isfinite(number) or not number.is_integer() or number < minimum:
        raise ValueError(f"{label}必须是不小于 {minimum} 的整数")
    return int(number)


def _parameters(raw):
    if raw is None or raw in ("", "None", "null"):
        return {}
    if not isinstance(raw, str) or len(raw) > 1_000_000:
        raise ValueError("旧版参数必须是长度合理的字典文本")
    try:
        value = json.loads(raw)
    except (ValueError, TypeError):
        try:
            value = ast.literal_eval(raw)
        except (ValueError, SyntaxError, RecursionError) as error:
            raise ValueError("旧版参数不是安全的字典文本（不会执行其中代码）") from error
    if not isinstance(value, dict) or not all(isinstance(key, str) for key in value):
        raise ValueError("旧版参数必须是以文字为键的字典")
    # Reject sets, bytes, NaN and other non-JSON values; tuples become arrays.
    return json.loads(json.dumps(value, ensure_ascii=False, allow_nan=False))


def _convert_parameters(type_id, payload, raw, base_directory):
    if type_id not in SUPPORTED_TYPES:
        raise ValueError(f"暂不能无损转换旧指令“{type_id}”，请提供原文件以补充适配")
    parameters = _parameters(raw)
    if type_id in PAYLOAD_FIELDS:
        key = PAYLOAD_FIELDS[type_id]
        if payload is not None:
            if key in parameters and parameters[key] != str(payload):
                raise ValueError(f"{key}在两个旧字段中内容冲突")
            parameters[key] = str(payload)
    elif payload not in (None, "", "None"):
        raise ValueError("该指令包含尚未识别的图像名称/内容字段")
    if type_id == "多图点击" and "图像路径" in parameters:
        parameters["图像路径"] = parameters["图像路径"].replace("、", "\n")
    if type_id == '坐标点击' and parameters.get('自定义次数') == 0 and parameters.get('动作') != '左键（自定义次数）':
        # Old UI stored zero even for single/double/move actions; it was unused.
        parameters['自定义次数'] = 1
    if type_id == "图像等待":
        from instructions.common.image_wait import normalize_image_wait
        parameters = normalize_image_wait(parameters)
    if type_id == "时间等待":
        for key in ("最小", "最大"):
            value = parameters.get(key)
            if isinstance(value, str) and "-" in value:
                parameters[key], parameters[key + "单位"] = value.rsplit("-", 1)
    if type_id == "滚轮滑动":
        parameters["方向"] = {"↑": "向上", "↓": "向下"}.get(
            parameters.get("方向"), parameters.get("方向", "向下")
        )
    if type_id == "鼠标拖拽" and "移动速度" in parameters:
        parameters["移动速度"] = float(parameters["移动速度"]) / 1000
    if type_id == "窗口焦点等待" and "检测频率" in parameters:
        parameters["检测频率"] = float(parameters["检测频率"]) / 1000
    if type_id == "窗口焦点等待":
        parameters['等待类型'] = {'等待窗口获取焦点': '获得焦点', '等待窗口失去焦点': '失去焦点'}.get(
            parameters.get('等待类型'), parameters.get('等待类型', '获得焦点'))
    if str(parameters.get("区域", "")).replace(" ", "") in {"(0,0,0,0)", "[0,0,0,0]"}:
        parameters["区域"] = ""

    fields = get_instruction_spec(type_id).load_editor_class().FIELDS
    unknown = set(parameters) - {field.key for field in fields}
    if unknown:
        raise ValueError(f"包含尚未适配的旧参数：{'、'.join(sorted(unknown))}")
    for field in fields:
        if field.key not in parameters:
            if field.required and field.default == "":
                raise ValueError(f"缺少必需参数：{field.label}")
            continue
        value = parameters[field.key]
        if field.kind == "bool":
            if isinstance(value, bool):
                pass
            elif str(value).strip().lower() in {"true", "false", "1", "0"}:
                value = str(value).strip().lower() in {"true", "1"}
            else:
                raise ValueError(f"{field.label}不是有效的真假值")
        elif field.kind in {"int", "float"}:
            if isinstance(value, bool):
                raise ValueError(f"{field.label}不是有效数字")
            value = float(value)
            if not math.isfinite(value) or not field.minimum <= value <= field.maximum:
                raise ValueError(f"{field.label}超出新版允许范围，不会截断数据")
            if field.kind == "int":
                value = _integer(value, field.label, field.minimum)
        elif field.kind == "choice" and value not in field.choices:
            raise ValueError(f"{field.label}的旧选项“{value}”尚未适配")
        elif field.kind in {"text", "multiline", "path"}:
            value = "" if value is None else str(value)
        if field.required and value in (None, ""):
            raise ValueError(f"必需参数不能为空：{field.label}")
        parameters[field.key] = value

    # Relative resources may be next to the old export. Resolve only existing
    # files, without guessing a replacement for missing resources or variables.
    if base_directory:
        for key in ("图像路径", "工作簿", "文件路径"):
            if key in parameters:
                paths = str(parameters[key]).splitlines()
                parameters[key] = "\n".join(
                    str((Path(base_directory) / path).resolve())
                    if path and not Path(path).is_absolute()
                    and (Path(base_directory) / path).is_file() else path
                    for path in paths
                )
    return parameters


def convert_legacy_workbook(source, base_directory=None, *, branch=None, branch_targets=None):
    """Return a new four-sheet workbook, or the unchanged current workbook.

    Fail closed for unknown sheets, auxiliary parameters and branch semantics.
    No source file, database, process, keyboard or mouse operation occurs here.
    """
    if is_current_workbook(source):
        return source
    target = Workbook()
    target.remove(target.active)
    for name, headers in zip(WORKBOOK_SHEETS, (
        COMMAND_SHEET_HEADERS, NODE_SHEET_HEADERS, EDGE_SHEET_HEADERS, SETTINGS_SHEET_HEADERS
    )):
        target.create_sheet(name).append(headers)
    active_sheets = []
    records = []
    jumps = []
    for sheet in source.worksheets:
        if sheet.title == "设置":
            continue
        if branch is not None and sheet.title != branch:
            continue
        headers = list(next(sheet.iter_rows(max_row=1, values_only=True)))
        if headers not in [LEGACY_HEADERS, DATABASE_HEADERS,
                           LEGACY_HEADERS + ["隶属分支"], DATABASE_HEADERS + ["隶属分支"]]:
            raise WorkbookValidationError(f"“{sheet.title}”不是可识别的旧版指令表；请提供原文件，当前数据未修改")
        populated = False
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if all(value is None for value in row):
                continue
            populated = True
            try:
                command_id, payload, type_id, raw, p2, p3, p4, repeat, policy, note = row[:10]
                if any(value not in (None, "", "None") for value in (p2, p3, p4)):
                    raise ValueError("参数 2–4 包含尚未适配的数据")
                if len(row) == 11 and row[10] not in (None, "", sheet.title):
                    raise ValueError("隶属分支与工作表名称不一致")
                policy = policy or "提示异常并暂停"
                if policy not in POLICIES:
                    match = re.fullmatch(r'(.+)-([1-9][0-9]*)', str(policy))
                    if not match:
                        raise ValueError(f"异常处理“{policy}”不是可识别的分支名-行号")
                    destination, target_row = match[1], int(match[2])
                    target_sheet = source[destination] if destination in source.sheetnames else None
                    target_count = sum(any(v is not None for v in r) for r in
                                       target_sheet.iter_rows(min_row=2, values_only=True)) if target_sheet else 0
                    if destination != sheet.title and branch_targets is not None:
                        target_count = branch_targets.get(destination, ('', 0))[1]
                    if not 1 <= target_row <= target_count:
                        raise ValueError(f'跳转目标“{policy}”不存在或行号越界')
                    if destination != sheet.title and (branch_targets is None or destination not in branch_targets):
                        raise ValueError('跨分支跳转必须使用完整迁移器生成关联项目，不能改为顺序执行')
                    jumps.append((len(records), destination, target_row))
                    policy = '提示异常并停止'
                parameters = _convert_parameters(str(type_id), payload, raw, base_directory)
                records.append([
                    _integer(command_id, "ID"), str(type_id),
                    json.dumps(parameters, ensure_ascii=False, allow_nan=False),
                    _integer(repeat if repeat is not None else 1, "重复次数"),
                    policy, str(note or ""), len(records),
                ])
            except (ValueError, TypeError, OverflowError, KeyError) as error:
                raise WorkbookValidationError(f"“{sheet.title}”第 {row_number} 行：{error}；当前数据未修改") from error
        if populated:
            active_sheets.append(sheet.title)
    if not records:
        raise WorkbookValidationError("旧工作簿没有可转换的指令，当前数据未修改")
    if len(active_sheets) != 1:
        raise WorkbookValidationError("旧文件含多个非空分支，请使用迁移器逐分支转换并选择要运行的分支；当前数据未修改")
    if len({record[0] for record in records}) != len(records):
        raise WorkbookValidationError("旧指令 ID 重复，当前数据未修改")
    if "设置" in source.sheetnames:
        settings = source["设置"]
        if list(next(settings.iter_rows(max_row=1, values_only=True))) != SETTINGS_SHEET_HEADERS:
            raise WorkbookValidationError("旧版设置表格式尚未识别，当前数据未修改")
        for row in settings.iter_rows(min_row=2, values_only=True):
            if all(value is None for value in row):
                continue
            if row[0] == "分支":
                if row[1] == active_sheets[0] and (row[2] not in (None, "") or row[3] not in (None, 1, "1")):
                    raise WorkbookValidationError("旧分支包含快捷键或整组重复设置，需要单独适配，当前数据未修改")
                continue
            target["设置"].append(row)
    target["节点"].append([START_NODE_ID, None, "start", 80, 80])
    previous = START_NODE_ID
    for index, record in enumerate(records):
        target["命令"].append(record)
        node_id = f"legacy-{record[0]}"
        target["节点"].append([node_id, record[0], "instruction", 80, 240 + index * 160])
        target["连线"].append([previous, node_id, 0])
        previous = node_id
    target["节点"].append([END_NODE_ID, None, "end", 80, 240 + len(records) * 160])
    target["连线"].append([previous, END_NODE_ID, 0])
    # Helpers are appended after original commands so legacy row numbers stay stable.
    helper_id = max(record[0] for record in records) + 1
    for offset, (source_index, destination, target_row) in enumerate(jumps):
        command_id = helper_id + offset
        node_id = f'legacy-error-{command_id}'
        same_branch = destination == active_sheets[0]
        parameters = {'跳转方式': '连线节点' if same_branch else '其他项目行',
                      '目标行': target_row,
                      '项目路径': '' if same_branch else branch_targets[destination][0]}
        if not same_branch:
            parameters['旧版转移后结束'] = True
        target['命令'].append([command_id, '报错跳转', json.dumps(parameters, ensure_ascii=False),
                               1, '提示异常并停止', f'旧版异常跳转：{destination}-{target_row}', len(records)+offset])
        target['节点'].append([node_id, command_id, 'instruction', 440, 240+source_index*160])
        target['连线'].append([f'legacy-{records[source_index][0]}', node_id, 5])
        if same_branch:
            target['连线'].append([node_id, f'legacy-{records[target_row-1][0]}', 6])
    return target


def save_converted_copy(workbook, source_path):
    """Reserve a sibling filename exclusively; never overwrite an old export."""
    source_path = Path(source_path)
    for index in range(10000):
        suffix = "" if index == 0 else f"-{index}"
        output = source_path.with_name(f"{source_path.stem}-新版{suffix}.xlsx")
        try:
            stream = output.open("xb")
        except FileExistsError:
            continue
        try:
            with stream:
                workbook.save(stream)
        except Exception:
            output.unlink(missing_ok=True)  # Only our just-created incomplete file.
            raise
        return str(output)
    raise OSError("无法生成不重名的新版文件名")
